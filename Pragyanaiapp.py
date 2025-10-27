import streamlit as st
import os
import pdfplumber
import docx
import openpyxl
import json
import tempfile
from groq import Groq
from gtts import gTTS
import traceback
import re
from dotenv import load_dotenv 

# -------------------------
# CONFIGURATION & API SETUP
# -------------------------

# CRITICAL FIX: Using the currently supported Groq model.
GROQ_MODEL = "llama-3.1-8b-instant"

# Options for LLM functions (defined for use in Candidate Dashboard)
section_options = ["name", "email", "phone", "skills", "education", "experience", "certifications", "projects", "strength", "personal_details", "github", "linkedin", "full resume"]
question_section_options = ["skills","experience", "certifications", "education", "projects"]
answer_types = [("Point-wise", "points"), ("Detailed", "detailed"), ("Key Points", "key")]


# Load environment variables from .env file
load_dotenv()

GROQ_API_KEY = os.getenv('GROQ_API_KEY')

if not GROQ_API_KEY:
    # Fail early with a clear message if the key is missing.
    st.error(
        "🚨 FATAL ERROR: GROQ_API_KEY environment variable not set. "
        "Please ensure a '.env' file exists in the script directory with this line: "
        "GROQ_API_KEY=\"YOUR_KEY_HERE\""
    )
    st.stop()

# Initialize Groq Client
client = Groq(api_key=GROQ_API_KEY)


# -------------------------
# Utility: Navigation Manager
# -------------------------
def go_to(page_name):
    """Changes the current page in Streamlit's session state."""
    st.session_state.page = page_name

# -------------------------
# CORE LOGIC: FILE HANDLING AND EXTRACTION
# -------------------------

def get_file_type(file_path):
    """Identifies the file type based on its extension."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pdf':
        return 'pdf'
    elif ext == '.docx':
        return 'docx'
    else:
        # Assuming other file types like .txt, .json are treated as plain text
        return 'txt' 

def extract_content(file_type, file_path):
    """Extracts text content from PDF or DOCX files using robust libraries."""
    try:
        if file_type == 'pdf':
            text = ''
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + '\n'
            if not text.strip():
                return "Error: PDF extraction failed. The file might be a scanned image without searchable text or is empty."
            return text
        
        elif file_type == 'docx':
            doc = docx.Document(file_path)
            text = '\n'.join([para.text for para in doc.paragraphs])
            if not text.strip():
                return "Error: DOCX content extraction failed. The file appears to be empty."
            return text
        
        elif file_type == 'txt':
            # Handle plain text, .json, or other unrecognized formats as text
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        
        else:
            return "Error: Unsupported file type."
    
    except Exception as e:
        return f"Fatal Extraction Error: Failed to read file content. Error details: {e}"

# -------------------------
# LLM & Extraction Functions
# -------------------------

@st.cache_data(show_spinner="Analyzing content with Groq LLM...")
def parse_with_llm(text, return_type='json'):
    """Sends resume text to the LLM for structured information extraction."""
    if text.startswith("Error"):
        return {"error": text, "raw_output": ""}

    prompt = f"""Extract the following information from the resume in structured JSON.
    Ensure all relevant details for each category are captured.
    - Name, - Email, - Phone, - Skills, - Education (list of degrees/institutions/dates), 
    - Experience (list of job roles/companies/dates/responsibilities), - Certifications (list), 
    - Projects (list of project names/descriptions/technologies), - Strength (list of personal strengths/qualities), 
    - Personal Details (e.g., address, date of birth, nationality), - Github (URL), - LinkedIn (URL)
    
    Resume Text:
    {text}
    
    Provide the output strictly as a JSON object.
    """
    content = ""
    try:
        response = client.chat.completions.create(
            model=GROQ_MODEL,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.2
        )
        content = response.choices[0].message.content.strip()

        # Robust JSON extraction
        json_str = content
        if json_str.startswith('```json'):
            json_str = json_str[len('```json'):]
        if json_str.endswith('```'):
            json_str = json_str[:-len('```')]
        json_str = json_str.strip()

        json_start = json_str.find('{')
        json_end = json_str.rfind('}') + 1

        if json_start != -1 and json_end != -1 and json_end > json_start:
            json_str = json_str[json_start:json_end]

        parsed = json.loads(json_str)

    except json.JSONDecodeError as e:
        error_msg = f"JSON decoding error from LLM. LLM returned malformed JSON. Error: {e}"
        parsed = {"error": error_msg, "raw_output": content}
    except Exception as e:
        error_msg = f"LLM API interaction error: {e}"
        parsed = {"error": error_msg, "raw_output": "No LLM response due to API error."}

    if return_type == 'json':
        return parsed
    elif return_type == 'markdown':
        if "error" in parsed:
            return f"**Error:** {parsed.get('error', 'Unknown parsing error')}\nRaw output:\n```\n{parsed.get('raw_output','')}\n```"
        
        md = ""
        for k, v in parsed.items():
            if v:
                md += f"**{k.replace('_', ' ').title()}**:\n"
                if isinstance(v, list):
                    for item in v:
                        if item: 
                            md += f"- {item}\n"
                elif isinstance(v, dict):
                    for sub_k, sub_v in v.items():
                        if sub_v:
                            md += f"  - {sub_k.replace('_', ' ').title()}: {sub_v}\n"
                else:
                    md += f"  {v}\n"
                md += "\n"
        return md
    return {"error": "Invalid return_type"}


# **REMOVED** the inaccurate extract_jd_from_linkedin_url function

def evaluate_jd_fit(job_description, parsed_json):
    """Evaluates how well a resume fits a given job description, including section-wise scores."""
    if not job_description.strip(): return "Please paste a job description."
    
    relevant_resume_data = {
        'Skills': parsed_json.get('skills', 'Not found or empty'),
        'Experience': parsed_json.get('experience', 'Not found or empty'),
        'Education': parsed_json.get('education', 'Not found or empty'),
    }
    resume_summary = json.dumps(relevant_resume_data, indent=2)

    prompt = f"""Evaluate how well the following resume content matches the provided job description.
    
    Job Description: {job_description}
    
    Resume Sections for Analysis:
    {resume_summary}
    
    Provide a detailed evaluation structured as follows:
    1.  **Overall Fit Score:** A score out of 10.
    2.  **Section Match Percentages:** A percentage score for the match in the key sections (Skills, Experience, Education).
    3.  **Strengths/Matches:** Key points where the resume aligns well with the JD.
    4.  **Gaps/Areas for Improvement:** Key requirements in the JD that are missing or weak in the resume.
    5.  **Overall Summary:** A concise summary of the fit.
    
    **Format the output strictly as follows:**
    Overall Fit Score: [Score]/10
    
    --- Section Match Analysis ---
    Skills Match: [XX]%
    Experience Match: [YY]%
    Education Match: [ZZ]%
    
    Strengths/Matches:
    - Point 1
    - Point 2
    
    Gaps/Areas for Improvement:
    - Point 1
    - Point 2
    
    Overall Summary: [Concise summary]
    """

    response = client.chat.completions.create(
        model=GROQ_MODEL, 
        messages=[{"role": "user", "content": prompt}], 
        temperature=0.3
    )
    return response.choices[0].message.content.strip()


# -------------------------
# Utility Functions
# -------------------------
def dump_to_excel(parsed_json, filename):
    """Dumps parsed JSON data to an Excel file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Profile Data"
    ws.append(["Category", "Details"])
    
    section_order = ['name', 'email', 'phone', 'github', 'linkedin', 'experience', 'education', 'skills', 'projects', 'certifications', 'strength', 'personal_details']
    
    for section_key in section_order:
        if section_key in parsed_json and parsed_json[section_key]:
            content = parsed_json[section_key]
            
            if section_key in ['name', 'email', 'phone', 'github', 'linkedin']:
                ws.append([section_key.replace('_', ' ').title(), str(content)])
            else:
                ws.append([])
                ws.append([section_key.replace('_', ' ').title()])
                
                if isinstance(content, list):
                    for item in content:
                        ws.append(["", str(item)])
                elif isinstance(content, dict):
                    for k, v in content.items():
                        ws.append(["", f"{k.replace('_', ' ').title()}: {v}"])
                else:
                    ws.append(["", str(content)])

    wb.save(filename)
    with open(filename, "rb") as f:
        return f.read()

def parse_and_store_resume(uploaded_file, file_name_key='default'):
    """Handles file upload, parsing, and stores results in session state."""
    
    temp_dir = tempfile.mkdtemp()
    temp_path = os.path.join(temp_dir, uploaded_file.name)
    with open(temp_path, "wb") as f:
        f.write(uploaded_file.getbuffer())

    file_type = get_file_type(temp_path)
    text = extract_content(file_type, temp_path)
    
    if text.startswith("Error"):
        return {"error": text, "full_text": text}

    parsed = parse_with_llm(text, return_type='json')
    
    if not parsed or "error" in parsed:
        return {"error": parsed.get('error', 'Unknown parsing error'), "full_text": text}

    # Generate Excel data for download if needed (only for single resume upload in Candidate dashboard)
    excel_data = None
    if file_name_key == 'single_resume_candidate':
        try:
            name = parsed.get('name', 'candidate').replace(' ', '_').strip()
            name = "".join(c for c in name if c.isalnum() or c in ('_', '-')).rstrip()
            if not name: name = "candidate"
            excel_filename = os.path.join(tempfile.gettempdir(), f"{name}_parsed_data.xlsx")
            excel_data = dump_to_excel(parsed, excel_filename)
        except Exception as e:
            st.warning(f"Could not generate Excel file for single resume: {e}")
    
    return {
        "parsed": parsed,
        "full_text": text,
        "excel_data": excel_data,
        "name": parsed.get('name', uploaded_file.name.split('.')[0])
    }


def qa_on_resume(question):
    """Chatbot for Resume (Q&A) using LLM."""
    parsed_json = st.session_state.parsed
    full_text = st.session_state.full_text
    prompt = f"""Given the following resume information:
    Resume Text: {full_text}
    Parsed Resume Data (JSON): {json.dumps(parsed_json, indent=2)}
    Answer the following question about the resume concisely and directly.
    If the information is not present, state that clearly.
    Question: {question}
    """
    response = client.chat.completions.create(model=GROQ_MODEL, messages=[{"role": "user", "content": prompt}], temperature=0.4)
    return response.choices[0].message.content.strip()

def generate_interview_questions(parsed_json, section):
    """Generates categorized interview questions using LLM."""
    section_title = section.replace("_", " ").title()
    section_content = parsed_json.get(section, "")
    if isinstance(section_content, (list, dict)):
        section_content = json.dumps(section_content, indent=2)
    elif not isinstance(section_content, str):
        section_content = str(section_content)

    if not section_content.strip():
        return f"No significant content found for the '{section_title}' section in the parsed resume. Please select a section with relevant data to generate questions."

    prompt = f"""Based on the following {section_title} section from the resume: {section_content}
Generate 3 interview questions each for these levels: Generic, Basic, Intermediate, Difficult.
**IMPORTANT: Format the output strictly as follows, with level headers and questions starting with 'Qx:':**
[Generic]
Q1: Question text...
...
[Difficult]
Q1: Question text...
    """
    response = client.chat.completions.create(
        model=GROQ_MODEL, 
        messages=[{"role": "user", "content": prompt}], 
        temperature=0.5
    )
    return response.choices[0].message.content.strip()


# -------------------------
# UI PAGES: Authentication
# -------------------------
def login_page():
    st.title("🌐 PragyanAI Job Portal")
    st.header("Login")

    email = st.text_input("Email")
    password = st.text_input("Password", type="password")

    if st.button("Login", use_container_width=True):
        if email and password:
            # Simulate successful login and go to role selection
            st.success("Login successful!")
            go_to("role_selection")
        else:
            st.error("Please enter both email and password")

    st.markdown("---")
    
    if st.button("Don't have an account? Sign up here"):
        go_to("signup")

def signup_page():
    st.header("Create an Account")
    email = st.text_input("Email")
    password = st.text_input("Password", type="password")
    confirm = st.text_input("Confirm Password", type="password")

    if st.button("Sign Up", use_container_width=True):
        if password == confirm and email:
            st.success("Signup successful! Please login.")
            go_to("login")
        else:
            st.error("Passwords do not match or email is empty")

    if st.button("Already have an account? Login here"):
        go_to("login")

def role_selection_page():
    st.header("Select Your Role")
    role = st.selectbox(
        "Choose a Dashboard",
        ["Select Role", "Admin Dashboard", "Candidate Dashboard", "Hiring Company Dashboard"]
    )

    col1, col2 = st.columns([2, 1])

    with col1:
        if st.button("Continue", use_container_width=True):
            if role == "Admin Dashboard":
                go_to("admin_dashboard")
            elif role == "Candidate Dashboard":
                go_to("candidate_dashboard")
            elif role == "Hiring Company Dashboard":
                go_to("hiring_dashboard")
            else:
                st.warning("Please select a role first")

    with col2:
        if st.button("⬅️ Go Back to Login"):
            go_to("login")

# -------------------------
# UI PAGES: Dashboards
# -------------------------

# The Admin dashboard has been updated with the robust regex extraction fix and JD URL fix
def admin_dashboard():
    st.header("🧑‍💼 Admin Dashboard")
    st.sidebar.button("⬅️ Go Back to Role Selection", on_click=go_to, args=("role_selection",))
    
    # Initialize Admin session state variables
    if "admin_jd_list" not in st.session_state:
        st.session_state.admin_jd_list = []
    if "resumes_to_analyze" not in st.session_state:
        st.session_state.resumes_to_analyze = []
    if "admin_match_results" not in st.session_state:
        st.session_state.admin_match_results = []
    
    tab_jd, tab_analysis = st.tabs(["📄 Job Description Management", "📊 Resume Analysis"])

    # --- TAB 1: JD Management ---
    with tab_jd:
        st.subheader("Add and Manage Job Descriptions (JD)")
        
        jd_type = st.radio("Select JD Type", ["Single JD", "Multiple JD"], key="jd_type_admin")
        st.markdown("### Add JD by:")
        
        # Options for adding JD 
        method = st.radio("Choose Method", ["Upload File", "Paste Text", "LinkedIn URL"], key="jd_add_method_admin") 

        # URL
        if method == "LinkedIn URL":
            url_list = st.text_area(
                "Enter LinkedIn URL(s) for reference (optional, for naming)" if jd_type == "Multiple JD" else "Enter LinkedIn URL for reference (optional, for naming)", key="url_list_admin"
            )
            # --- START FIX: Capture Actual JD Text ---
            actual_jd_text = st.text_area(
                "Paste the **ACTUAL Job Description Text** copied from the LinkedIn page here (REQUIRED for accurate matching).", 
                key="actual_jd_text_admin",
                height=250
            )
            # --- END FIX: Capture Actual JD Text ---
            
            if st.button("Add JD(s) from URL", key="add_jd_url_btn_admin"):
                if not actual_jd_text:
                    st.error("Please paste the actual JD content for matching.")
                    return

                # Handle single JD case
                if jd_type == "Single JD":
                    url = url_list.split(",")[0].strip() if url_list else ""
                    name_base = url.split('/jobs/view/')[-1].split('/')[0] if '/jobs/view/' in url else "LinkedIn Job"
                    jd_name = f"JD from URL: {name_base}"
                    
                    st.session_state.admin_jd_list.append({"name": jd_name, "content": actual_jd_text})
                    st.success(f"✅ JD '{jd_name}' added successfully using the pasted text!")
                    
                # Handle multiple JD case - disallow due to single text area
                elif jd_type == "Multiple JD":
                    st.error("For Multiple JDs via URL, please switch to the 'Paste Text' method and separate JDs with '---', as we cannot map multiple URLs to a single text box.")


        # Paste Text
        elif method == "Paste Text":
            text_list = st.text_area(
                "Paste one or more JD texts (separate by '---')" if jd_type == "Multiple JD" else "Paste JD text here", key="text_list_admin"
            )
            if st.button("Add JD(s) from Text", key="add_jd_text_btn_admin"):
                if text_list:
                    texts = [t.strip() for t in text_list.split("---")] if jd_type == "Multiple JD" else [text_list.strip()]
                    for i, text in enumerate(texts):
                         if text:
                            # Use the first line as a name
                            name_base = text.splitlines()[0].strip()
                            if len(name_base) > 30: name_base = f"{name_base[:27]}..."
                            if not name_base: name_base = f"Pasted JD {len(st.session_state.admin_jd_list) + i + 1}"
                            
                            st.session_state.admin_jd_list.append({"name": name_base, "content": text})
                    st.success(f"✅ {len(texts)} JD(s) added successfully!")

        # Upload File
        elif method == "Upload File":
            uploaded_files = st.file_uploader(
                "Upload JD file(s)",
                type=["pdf", "txt", "docx"],
                accept_multiple_files=True if jd_type == "Multiple JD" else False,
                key="jd_file_uploader_admin"
            )
            if st.button("Add JD(s) from File", key="add_jd_file_btn_admin"):
                files_to_process = uploaded_files if jd_type == "Multiple JD" and uploaded_files else [uploaded_files]
                count = 0
                for file in files_to_process:
                    if file:
                        temp_dir = tempfile.mkdtemp()
                        temp_path = os.path.join(temp_dir, file.name)
                        with open(temp_path, "wb") as f:
                            f.write(file.getbuffer())
                            
                        file_type = get_file_type(temp_path)
                        jd_text = extract_content(file_type, temp_path)
                        
                        if not jd_text.startswith("Error"):
                            st.session_state.admin_jd_list.append({"name": file.name, "content": jd_text})
                            count += 1
                if count > 0:
                    st.success(f"✅ {count} JD(s) added successfully!")
                else:
                    st.error("No valid JD files were uploaded or content extraction failed.")

        # Display Added JDs
        if st.session_state.admin_jd_list:
            st.markdown("### ✅ Current JDs Added:")
            for idx, jd_item in enumerate(st.session_state.admin_jd_list, 1):
                title = jd_item['name']
                # Clean up simulated title prefix for display
                display_title = title.replace("--- Simulated JD for: ", "")
                with st.expander(f"JD {idx}: {display_title}"):
                    st.text(jd_item['content'])
        else:
            st.info("No Job Descriptions added yet.")


    # --- TAB 2: Resume Analysis (Admin logic) ---
    with tab_analysis:
        st.subheader("Analyze Resumes Against Job Descriptions")

        # 1. Resume Upload (Admin uses st.session_state.resumes_to_analyze list)
        st.markdown("#### 1. Upload Resumes")
        resume_upload_type = st.radio("Upload Type", ["Single Resume", "Multiple Resumes"], key="resume_upload_type_admin")

        uploaded_files = st.file_uploader(
            "Choose files to analyze",
            # Added more types for robustness
            type=["pdf", "docx", "txt", "json", "rtf"], 
            accept_multiple_files=(resume_upload_type == "Multiple Resumes"),
            key="resume_file_uploader_admin"
        )
        
        if st.button("Load and Parse Resume(s) for Analysis", key="parse_resumes_admin"):
            if uploaded_files:
                files_to_process = uploaded_files if isinstance(uploaded_files, list) else [uploaded_files]
                st.session_state.resumes_to_analyze = []
                count = 0
                with st.spinner("Parsing resume(s)... This may take a moment."):
                    for file in files_to_process:
                        if file:
                            # Use parse_and_store_resume logic, but store results in a list
                            result = parse_and_store_resume(file, file_name_key='admin_analysis')
                            
                            if "error" not in result:
                                st.session_state.resumes_to_analyze.append(result)
                                count += 1
                            else:
                                st.error(f"Failed to parse {file.name}: {result['error']}")

                if count > 0:
                    st.success(f"Successfully loaded and parsed {count} resume(s) for analysis.")
                elif not st.session_state.resumes_to_analyze:
                    st.warning("No resumes were successfully loaded and parsed.")
            else:
                st.warning("Please upload one or more resume files.")

        st.markdown("---")

        # 2. JD Selection and Analysis
        st.markdown("#### 2. Select JD and Run Analysis")

        if not st.session_state.resumes_to_analyze:
            st.info("Upload and parse resumes first to enable analysis.")
            return

        if not st.session_state.admin_jd_list:
            st.error("Please add at least one Job Description in the 'JD Management' tab before running an analysis.")
            return

        jd_options = {item['name']: item['content'] for item in st.session_state.admin_jd_list}
        selected_jd_name = st.selectbox("Select JD for Matching", list(jd_options.keys()), key="select_jd_admin")
        selected_jd_content = jd_options.get(selected_jd_name, "")


        if st.button("Run Match Analysis", key="run_match_analysis_admin"):
            st.session_state.admin_match_results = []
            
            if not selected_jd_content:
                st.error("Selected JD content is empty.")
                return

            with st.spinner(f"Matching {len(st.session_state.resumes_to_analyze)} resumes against '{selected_jd_name}'..."):
                for resume_data in st.session_state.resumes_to_analyze:
                    
                    resume_name = resume_data['name']
                    parsed_json = resume_data['parsed']

                    try:
                        fit_output = evaluate_jd_fit(selected_jd_content, parsed_json)
                        
                        # --- ENHANCED EXTRACTION LOGIC (FIXED) ---
                        # 1. Overall Score: Look for (number)/10, robust to newlines or extra spaces
                        overall_score_match = re.search(r'Overall Fit Score:\s*(\d+)\s*/10', fit_output, re.IGNORECASE)
                        
                        # 2. Section Matches: Look for "Key Match: XX%" pattern within the Section Analysis text
                        section_analysis_match = re.search(
                             r'--- Section Match Analysis ---\s*(.*?)\s*Strengths/Matches:', 
                             fit_output, re.DOTALL
                        )

                        skills_percent = 'N/A'
                        experience_percent = 'N/A'
                        education_percent = 'N/A'
                        
                        if section_analysis_match:
                            section_text = section_analysis_match.group(1)
                            
                            # Extract percentages from the collected section text
                            skills_match = re.search(r'Skills Match:\s*(\d+)%', section_text, re.IGNORECASE)
                            experience_match = re.search(r'Experience Match:\s*(\d+)%', section_text, re.IGNORECASE)
                            education_match = re.search(r'Education Match:\s*(\d+)%', section_text, re.IGNORECASE)
                            
                            if skills_match:
                                skills_percent = skills_match.group(1)
                            if experience_match:
                                experience_percent = experience_match.group(1)
                            if education_match:
                                education_percent = education_match.group(1)
                        
                        overall_score = overall_score_match.group(1) if overall_score_match else 'N/A'
                        # --- END ENHANCED EXTRACTION LOGIC (FIXED) ---

                        st.session_state.admin_match_results.append({
                            "resume_name": resume_name,
                            "jd_name": selected_jd_name,
                            "overall_score": overall_score,
                            "skills_percent": skills_percent,
                            "experience_percent": experience_percent, 
                            "education_percent": education_percent,   
                            "full_analysis": fit_output
                        })
                    except Exception as e:
                        st.session_state.admin_match_results.append({
                            "resume_name": resume_name,
                            "jd_name": selected_jd_name,
                            "overall_score": "Error",
                            "skills_percent": "Error",
                            "experience_percent": "Error", 
                            "education_percent": "Error",   
                            "full_analysis": f"Error running analysis: {e}\n{traceback.format_exc()}"
                        })
                st.success("Analysis complete!")


        # 3. Display Results
        if st.session_state.get('admin_match_results'):
            st.markdown("#### 3. Match Results")
            results_df = st.session_state.admin_match_results
            
            # Create a simple table/summary of results
            display_data = []
            for item in results_df:
                display_data.append({
                    "Resume": item["resume_name"],
                    "JD": item["jd_name"],
                    "Fit Score (out of 10)": item["overall_score"],
                    "Skills (%)": item.get("skills_percent", "N/A"),
                    "Experience (%)": item.get("experience_percent", "N/A"), 
                    "Education (%)": item.get("education_percent", "N/A"),   
                })

            st.dataframe(display_data, use_container_width=True)

            # Display detailed analysis in expanders
            st.markdown("##### Detailed Reports")
            for item in results_df:
                header_text = f"Report for **{item['resume_name']}** against {item['jd_name']} (Score: **{item['overall_score']}/10** | S: **{item.get('skills_percent', 'N/A')}%** | E: **{item.get('experience_percent', 'N/A')}%** | Edu: **{item.get('education_percent', 'N/A')}%**)"
                with st.expander(header_text):
                    st.markdown(item['full_analysis'])

# Candidate Dashboard is updated here
def candidate_dashboard():
    st.header("👩‍🎓 Candidate Dashboard")
    st.markdown("Welcome! Use the tabs below to upload your resume and access AI preparation tools.")

    st.sidebar.button("⬅️ Go Back to Role Selection", on_click=go_to, args=("role_selection",))
    
    # Sidebar for Resume Upload (Centralized Upload)
    with st.sidebar:
        st.header("Upload Your Resume")
        uploaded_file = st.file_uploader("Choose a PDF or DOCX file", type=["pdf", "docx"])
        
        if uploaded_file is not None:
            if st.button("Parse Resume", use_container_width=True):
                # Centralized upload logic for Candidate Dashboard
                result = parse_and_store_resume(uploaded_file, file_name_key='single_resume_candidate')
                
                if "error" not in result:
                    st.session_state.parsed = result['parsed']
                    st.session_state.full_text = result['full_text']
                    st.session_state.excel_data = result['excel_data'] # This must be set here
                    st.success("Resume parsed successfully!")
                else:
                    st.error(f"Parsing failed: {result['error']}")

        
        st.markdown("---")
        if st.session_state.parsed.get("name"):
            st.success(f"Resume for **{st.session_state.parsed['name']}** is loaded.")
        elif st.session_state.full_text:
            st.warning("Resume file loaded, but parsing may have errors.")
        else:
            st.info("Please upload a resume to begin.")

    # Main Content Tabs (AI Resume Parser Features)
    tab1, tab2, tab3, tab4, tab5 = st.tabs([
        "📄 Resume Parsing", 
        "💬 Resume Chatbot (Q&A)", 
        "❓ Interview Prep", 
        "📚 JD Management",
        "🎯 Batch JD Match"
    ])
    
    # --- TAB 1: Resume Parsing ---
    with tab1:
        st.header("Resume Parsing")
        if not st.session_state.full_text:
            st.warning("Please upload and parse a resume in the sidebar first.")
            return

        col1, col2 = st.columns(2)
        with col1:
            output_format = st.radio('Output Format', ['json', 'markdown'], key='format_radio_c')
        with col2:
            section = st.selectbox('Select Section to View', section_options, key='section_select_c')

        parsed = st.session_state.parsed
        full_text = st.session_state.full_text

        if "error" in parsed:
            st.error(parsed.get("error", "An unknown error occurred during parsing."))
            return

        # Display Main Parsed Output
        if output_format == 'json':
            output_str = json.dumps(parsed, indent=2)
            st.text_area("Parsed Output (JSON)", output_str, height=350)
        else:
            output_str = parse_with_llm(full_text, return_type='markdown')
            st.markdown("### Parsed Output (Markdown)")
            st.markdown(output_str)

        # Download Buttons
        if st.session_state.excel_data:
            st.download_button(
                label="Download Parsed Data (Excel)",
                data=st.session_state.excel_data,
                file_name=f"{parsed.get('name', 'candidate').replace(' ', '_')}_parsed_data.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        
        # Section View
        section_content_str = ""
        if section == "full resume":
            section_content_str = full_text
        elif section in parsed:
            section_val = parsed[section]
            section_content_str = json.dumps(section_val, indent=2) if isinstance(section_val, (list, dict)) else str(section_val)
        else:
            section_content_str = f"Section '{section}' not found or is empty."

        st.text_area("Selected Section Content", section_content_str, height=200)

    # --- TAB 2: Resume Chatbot (Q&A) ---
    with tab2:
        st.header("Resume Chatbot (Q&A)")
        st.markdown("### Ask any question about the uploaded resume.")
        if not st.session_state.full_text:
            st.warning("Please upload and parse a resume first.")
            return

        question = st.text_input("Your Question", placeholder="e.g., What are the candidate's key skills?")
        
        if st.button("Get Answer", key="qa_btn"):
            with st.spinner("Generating answer..."):
                try:
                    answer = qa_on_resume(question)
                    st.session_state.qa_answer = answer
                except Exception as e:
                    st.error(f"Error during Q&A: {e}")
                    st.session_state.qa_answer = "Could not generate an answer."

        if st.session_state.get('qa_answer'):
            st.text_area("Answer", st.session_state.qa_answer, height=150)

    # --- TAB 3: Interview Prep ---
    with tab3:
        st.header("Interview Preparation Tools")
        if not st.session_state.parsed or "error" in st.session_state.parsed:
            st.warning("Please upload and successfully parse a resume first.")
            return

        st.subheader("Generate Interview Questions")
        section_choice = st.selectbox("Select Section", question_section_options, key='iq_section_c')
        
        if st.button("Generate Interview Questions", key='iq_btn_c'):
            with st.spinner("Generating questions..."):
                try:
                    raw_questions_response = generate_interview_questions(st.session_state.parsed, section_choice)
                    st.session_state.iq_output = raw_questions_response
                except Exception as e:
                    st.error(f"Error generating questions: {e}")
                    st.session_state.iq_output = "Error generating questions."

        if st.session_state.get('iq_output'):
            st.text_area("Generated Interview Questions (by difficulty level)", st.session_state.iq_output, height=400)
            
    # --- TAB 4: JD Management (Modified Admin Logic) ---
    with tab4:
        st.header("📚 Manage Job Descriptions for Matching")
        st.markdown("Add multiple JDs here to compare your resume against them in the next tab.")
        
        # Initialize JD list specific to candidate if not present (to avoid mixing with admin's list)
        if "candidate_jd_list" not in st.session_state:
             st.session_state.candidate_jd_list = []
        
        jd_type = st.radio("Select JD Type", ["Single JD", "Multiple JD"], key="jd_type_candidate")
        st.markdown("### Add JD by:")
        
        # Options for adding JD 
        method = st.radio("Choose Method", ["Upload File", "Paste Text", "LinkedIn URL"], key="jd_add_method_candidate") 

        # URL
        if method == "LinkedIn URL":
            url_list = st.text_area(
                "Enter LinkedIn URL for reference (optional, for naming)", key="url_list_candidate"
            )
            # --- START FIX: Capture Actual JD Text ---
            actual_jd_text = st.text_area(
                "Paste the **ACTUAL Job Description Text** copied from the LinkedIn page here (REQUIRED for accurate matching).", 
                key="actual_jd_text_candidate",
                height=250
            )
            # --- END FIX: Capture Actual JD Text ---

            if st.button("Add JD(s) from URL", key="add_jd_url_btn_candidate"):
                if not actual_jd_text:
                    st.error("Please paste the actual JD content for matching.")
                    return

                # Handle single JD case
                if jd_type == "Single JD":
                    url = url_list.split(",")[0].strip() if url_list else ""
                    name_base = url.split('/jobs/view/')[-1].split('/')[0] if '/jobs/view/' in url else "LinkedIn Job"
                    jd_name = f"JD from URL: {name_base}"
                    
                    st.session_state.candidate_jd_list.append({"name": jd_name, "content": actual_jd_text})
                    st.success(f"✅ JD '{jd_name}' added successfully using the pasted text!")
                
                # Handle multiple JD case - disallow due to single text area
                elif jd_type == "Multiple JD":
                    st.error("For Multiple JDs via URL, please switch to the 'Paste Text' method and separate JDs with '---', as we cannot map multiple URLs to a single text box.")


        # Paste Text
        elif method == "Paste Text":
            text_list = st.text_area(
                "Paste one or more JD texts (separate by '---')" if jd_type == "Multiple JD" else "Paste JD text here", key="text_list_candidate"
            )
            if st.button("Add JD(s) from Text", key="add_jd_text_btn_candidate"):
                if text_list:
                    texts = [t.strip() for t in text_list.split("---")] if jd_type == "Multiple JD" else [text_list.strip()]
                    for i, text in enumerate(texts):
                         if text:
                            name_base = text.splitlines()[0].strip()
                            if len(name_base) > 30: name_base = f"{name_base[:27]}..."
                            if not name_base: name_base = f"Pasted JD {len(st.session_state.candidate_jd_list) + i + 1}"
                            
                            st.session_state.candidate_jd_list.append({"name": name_base, "content": text})
                    st.success(f"✅ {len(texts)} JD(s) added successfully!")

        # Upload File
        elif method == "Upload File":
            uploaded_files = st.file_uploader(
                "Upload JD file(s)",
                type=["pdf", "txt", "docx"],
                accept_multiple_files=True if jd_type == "Multiple JD" else False,
                key="jd_file_uploader_candidate"
            )
            if st.button("Add JD(s) from File", key="add_jd_file_btn_candidate"):
                files_to_process = uploaded_files if jd_type == "Multiple JD" and uploaded_files else [uploaded_files]
                count = 0
                for file in files_to_process:
                    if file:
                        temp_dir = tempfile.mkdtemp()
                        temp_path = os.path.join(temp_dir, file.name)
                        with open(temp_path, "wb") as f:
                            f.write(file.getbuffer())
                            
                        file_type = get_file_type(temp_path)
                        jd_text = extract_content(file_type, temp_path)
                        
                        if not jd_text.startswith("Error"):
                            st.session_state.candidate_jd_list.append({"name": file.name, "content": jd_text})
                            count += 1
                if count > 0:
                    st.success(f"✅ {count} JD(s) added successfully!")
                else:
                    st.error("No valid JD files were uploaded or content extraction failed.")

        # Display Added JDs
        if st.session_state.candidate_jd_list:
            st.markdown("### ✅ Current JDs Added:")
            for idx, jd_item in enumerate(st.session_state.candidate_jd_list, 1):
                title = jd_item['name']
                display_title = title.replace("--- Simulated JD for: ", "")
                with st.expander(f"JD {idx}: {display_title}"):
                    st.text(jd_item['content'])
        else:
            st.info("No Job Descriptions added yet.")

    # --- TAB 5: Batch JD Match (Modified Admin Analysis Logic) ---
    with tab5:
        st.header("🎯 Batch JD Match")
        st.markdown("Compare your current resume against all saved job descriptions.")

        if not st.session_state.parsed:
            st.warning("Please **upload and parse your resume** in the sidebar first.")
            return

        if not st.session_state.candidate_jd_list:
            st.error("Please **add Job Descriptions** in the 'JD Management' tab (Tab 4) before running batch analysis.")
            return
            
        # Initialize results list for the candidate dashboard
        if "candidate_match_results" not in st.session_state:
            st.session_state.candidate_match_results = []

        if st.button(f"Run Batch Match Against {len(st.session_state.candidate_jd_list)} JDs"):
            st.session_state.candidate_match_results = []
            
            resume_name = st.session_state.parsed.get('name', 'Uploaded Resume')
            parsed_json = st.session_state.parsed

            with st.spinner(f"Matching {resume_name}'s resume against {len(st.session_state.candidate_jd_list)} JDs..."):
                for jd_item in st.session_state.candidate_jd_list:
                    
                    jd_name = jd_item['name']
                    jd_content = jd_item['content']

                    try:
                        fit_output = evaluate_jd_fit(jd_content, parsed_json)
                        
                        # --- ENHANCED EXTRACTION LOGIC (FIXED) ---
                        # 1. Overall Score: Look for (number)/10, robust to newlines or extra spaces
                        overall_score_match = re.search(r'Overall Fit Score:\s*(\d+)\s*/10', fit_output, re.IGNORECASE)
                        
                        # 2. Section Matches: Look for "Key Match: XX%" pattern within the Section Analysis text
                        section_analysis_match = re.search(
                             r'--- Section Match Analysis ---\s*(.*?)\s*Strengths/Matches:', 
                             fit_output, re.DOTALL
                        )

                        skills_percent = 'N/A'
                        experience_percent = 'N/A'
                        education_percent = 'N/A'
                        
                        if section_analysis_match:
                            # Search only within the captured section block for better accuracy
                            section_text = section_analysis_match.group(1)
                            
                            # Extract percentages from the collected section text
                            skills_match = re.search(r'Skills Match:\s*(\d+)%', section_text, re.IGNORECASE)
                            experience_match = re.search(r'Experience Match:\s*(\d+)%', section_text, re.IGNORECASE)
                            education_match = re.search(r'Education Match:\s*(\d+)%', section_text, re.IGNORECASE)
                            
                            if skills_match:
                                skills_percent = skills_match.group(1)
                            if experience_match:
                                experience_percent = experience_match.group(1)
                            if education_match:
                                education_percent = education_match.group(1)
                        
                        overall_score = overall_score_match.group(1) if overall_score_match else 'N/A'
                        # --- END ENHANCED EXTRACTION LOGIC (FIXED) ---

                        st.session_state.candidate_match_results.append({
                            "jd_name": jd_name,
                            "overall_score": overall_score,
                            "skills_percent": skills_percent,
                            "experience_percent": experience_percent, 
                            "education_percent": education_percent,   
                            "full_analysis": fit_output
                        })
                    except Exception as e:
                        st.session_state.candidate_match_results.append({
                            "jd_name": jd_name,
                            "overall_score": "Error",
                            "skills_percent": "Error",
                            "experience_percent": "Error", 
                            "education_percent": "Error",   
                            "full_analysis": f"Error running analysis for {jd_name}: {e}\n{traceback.format_exc()}"
                        })
                st.success("Batch analysis complete!")


        # 3. Display Results
        if st.session_state.get('candidate_match_results'):
            st.markdown("#### Match Results for Your Resume")
            results_df = st.session_state.candidate_match_results
            
            # Create a simple table/summary of results
            display_data = []
            for item in results_df:
                display_data.append({
                    "Job Description": item["jd_name"].replace("--- Simulated JD for: ", ""),
                    "Fit Score (out of 10)": item["overall_score"],
                    "Skills (%)": item.get("skills_percent", "N/A"),
                    "Experience (%)": item.get("experience_percent", "N/A"), 
                    "Education (%)": item.get("education_percent", "N/A"),   
                })

            st.dataframe(display_data, use_container_width=True)

            # Display detailed analysis in expanders
            st.markdown("##### Detailed Reports")
            for item in results_df:
                header_text = f"Report for **{item['jd_name'].replace('--- Simulated JD for: ', '')}** (Score: **{item['overall_score']}/10** | S: **{item.get('skills_percent', 'N/A')}%** | E: **{item.get('experience_percent', 'N/A')}%** | Edu: **{item.get('education_percent', 'N/A')}%**)"
                with st.expander(header_text):
                    st.markdown(item['full_analysis'])


def hiring_dashboard():
    st.header("🏢 Hiring Company Dashboard")
    st.write("Manage job postings and view candidate applications. (Placeholder for future features)")
    st.sidebar.button("⬅️ Go Back to Role Selection", on_click=go_to, args=("role_selection",))

# -------------------------
# Main App Initialization
# -------------------------
def main():
    st.set_page_config(layout="wide", page_title="PragyanAI Job Portal")

    # --- Session State Initialization ---
    if 'page' not in st.session_state:
        st.session_state.page = "login"
    
    # Initialize session state for AI features
    if 'parsed' not in st.session_state:
        st.session_state.parsed = {}
        st.session_state.full_text = ""
        st.session_state.excel_data = None
        st.session_state.qa_answer = ""
        st.session_state.iq_output = ""
        st.session_state.jd_fit_output = ""
        
        # Admin Dashboard specific lists
        st.session_state.admin_jd_list = [] # For Admin Dashboard (list of dicts: {'name', 'content'})
        st.session_state.resumes_to_analyze = [] # For Admin Dashboard (list of dicts: {'name', 'parsed', 'full_text'})
        st.session_state.admin_match_results = [] # For Admin Dashboard match results
        
        # Candidate Dashboard specific lists
        st.session_state.candidate_jd_list = []
        st.session_state.candidate_match_results = []


    # --- Page Routing ---
    if st.session_state.page == "login":
        login_page()
    elif st.session_state.page == "signup":
        signup_page()
    elif st.session_state.page == "role_selection":
        role_selection_page()
    elif st.session_state.page == "admin_dashboard":
        admin_dashboard()
    elif st.session_state.page == "candidate_dashboard":
        candidate_dashboard()
    elif st.session_state.page == "hiring_dashboard":
        hiring_dashboard()

if __name__ == '__main__':
    main()
